from pathlib import Path
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
NAVY = "12304A"
TEAL = "008A8A"
GOLD = "F5B942"
LIGHT = "E8F1F0"
INK = "17212B"


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc.get_or_add_tcPr()
    tc_mar = tc.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc.append(tc_mar)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    marker = OxmlElement("w:tblHeader")
    marker.set(qn("w:val"), "true")
    tr_pr.append(marker)


def set_font(run, name="Arial", size=10.5, bold=False, color=INK, italic=False):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def configure(doc, title):
    doc.core_properties.title = title
    doc.core_properties.author = "Prof. Madya Dr. Mohd Murtadha Mohamad"
    sec = doc.sections[0]
    sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
    sec.top_margin = sec.bottom_margin = Cm(1.8)
    sec.left_margin = sec.right_margin = Cm(1.9)
    styles = doc.styles
    for name, size, color, bold in (("Normal", 10.5, INK, False), ("Heading 1", 17, NAVY, True), ("Heading 2", 13, TEAL, True), ("Heading 3", 11, NAVY, True)):
        st = styles[name]
        st.font.name = "Arial"
        st._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        st._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        st.font.size = Pt(size)
        st.font.bold = bold
        st.font.color.rgb = RGBColor.from_string(color)
        st.paragraph_format.space_after = Pt(5)
        st.paragraph_format.space_before = Pt(8 if name != "Normal" else 0)
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = footer.add_run("Bengkel Transformasi Digital MPE  |  28–29 Julai 2026")
    set_font(r, size=8, color="53636F")


def title_block(doc, title, subtitle):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(20)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(title)
    set_font(r, size=24, bold=True, color=NAVY)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(18)
    r = p.add_run(subtitle)
    set_font(r, size=11, italic=True, color=TEAL)


def bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(text)
    set_font(r)
    return p


def table(doc, headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    t.style = "Table Grid"
    repeat_header(t.rows[0])
    for i, text in enumerate(headers):
        c = t.rows[0].cells[i]
        shade(c, NAVY)
        margins(c)
        c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(text)
        set_font(r, size=9, bold=True, color="FFFFFF")
        if widths:
            c.width = Cm(widths[i])
    for ridx, row in enumerate(rows):
        cells = t.add_row().cells
        for i, text in enumerate(row):
            c = cells[i]
            if ridx % 2:
                shade(c, "F6F8F8")
            margins(c)
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = c.paragraphs[0]
            r = p.add_run(text)
            set_font(r, size=8.6)
            if widths:
                c.width = Cm(widths[i])
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return t


SCHEDULE = [
    ("28 Jul, 9.00–10.30", "Modul 1", "Aplikasi pintar untuk produktiviti MPE", "Gemini/ChatGPT, carian", "Peta tugasan + prompt"),
    ("28 Jul, 11.00–1.00", "Modul 2", "Automasi penyediaan dokumen rasmi", "Word/Docs, AI", "Draf + semak kelulusan"),
    ("28 Jul, 2.30–4.30", "Modul 3A", "Asas analisis data dan rekod teknikal", "Excel/Sheets", "Data bersih + daftar rekod"),
    ("29 Jul, 9.00–10.30", "Modul 3B", "Kualiti data dan tadbir urus rekod", "Excel/Sheets", "Audit data + prosedur"),
    ("29 Jul, 11.00–1.00", "Modul 4", "Aliran kerja dan pendigitalan proses makmal", "Forms, Automate/Script", "Peta proses + prototaip"),
    ("29 Jul, 2.30–4.30", "Modul 5", "Dokumen projek dan pembentangan", "Word/Docs, Slides", "Ringkasan + 3–5 slaid"),
]

MODULES = [
    ("Modul 1 — Aplikasi pintar untuk produktiviti MPE", "90 minit", "Gemini atau ChatGPT; enjin carian; prompt pack", ["Pencetus tugasan berulang — 10 minit", "Konsep dan demonstrasi Meeting Action Assistant — 20 minit", "Reka prompt bersempadan — 25 minit", "Uji kes biasa dan kes maklumat hilang — 25 minit", "Perkongsian manfaat dan risiko — 10 minit"], "Peta tugasan dan prompt produktiviti boleh guna semula."),
    ("Modul 2 — Automasi penyediaan dokumen rasmi", "120 minit", "Microsoft Word atau Google Docs; Gemini/ChatGPT; templat rasmi", ["Anatomi dan fakta dokumen — 15 minit", "Demonstrasi draf daripada fakta diluluskan — 20 minit", "Draf berpandu — 35 minit", "Semakan fakta dan format — 25 minit", "Semakan rakan dan pembaikan — 20 minit", "Rumusan kelulusan manusia — 5 minit"], "Draf dokumen rasmi dan senarai semak kelulusan."),
    ("Modul 3A — Asas analisis data dan rekod teknikal", "120 minit", "Microsoft Excel atau Google Sheets; AI untuk bantuan formula", ["Soalan analisis — 10 minit", "Import dan profil data — 20 minit", "Pembersihan data — 35 minit", "Analisis dan visual — 35 minit", "Daftar rekod teknikal — 15 minit", "Simpan dan refleksi — 5 minit"], "Set data bersih, ringkasan analisis dan daftar rekod."),
    ("Modul 3B — Kualiti data dan tadbir urus rekod", "90 minit", "Microsoft Excel atau Google Sheets", ["Imbas semula hasil Hari Pertama — 10 minit", "Audit lengkap, unik, sah, konsisten dan terkini — 30 minit", "Pembetulan serta jejak audit — 20 minit", "Tafsiran terkawal — 20 minit", "Prosedur rekod ringkas — 10 minit"], "Laporan kualiti data dan prosedur pengurusan rekod."),
    ("Modul 4 — Aliran kerja digital makmal", "120 minit", "Microsoft/Google Forms; Excel/Sheets; Power Automate atau Apps Script", ["Pilih proses — 10 minit", "Peta keadaan semasa — 25 minit", "Reka keadaan masa hadapan — 25 minit", "Bina prototaip — 35 minit", "Uji laluan biasa, tidak lengkap dan eskalasi — 20 minit", "Rumusan — 5 minit"], "Peta proses masa hadapan dan prototaip aliran kerja."),
    ("Modul 5 — Dokumen projek dan pembentangan", "120 minit", "Word/Docs; PowerPoint/Slides; Gemini/ChatGPT", ["Pilih hasil projek — 10 minit", "Susun konteks, masalah, bukti, cadangan dan tindakan — 20 minit", "Ringkasan satu halaman — 25 minit", "Bina 3–5 slaid — 30 minit", "Semakan rakan — 15 minit", "Pembentangan mini — 15 minit", "Pelan tindakan 30 hari — 5 minit"], "Ringkasan projek, pembentangan mini dan tindakan 30 hari."),
]


def build_facilitator():
    doc = Document()
    configure(doc, "Panduan Fasilitator — Bengkel Transformasi Digital MPE")
    title_block(doc, "Panduan Fasilitator", "Transformasi Digital dan Pemerkasaan Operasi Pintar Makmal Penyelidikan Elektrik (MPE)\n28–29 Julai 2026 · Prof. Madya Dr. Mohd Murtadha Mohamad")
    doc.add_heading("Tujuan penyampaian", 1)
    doc.add_paragraph("Membimbing peserta menghasilkan hasil kerja digital yang boleh disemak: prompt produktiviti, dokumen rasmi, analisis data, rekod teknikal, prototaip aliran kerja dan pembentangan projek.")
    doc.add_heading("Jadual modul", 1)
    table(doc, ["Masa", "Modul", "Fokus", "Alat", "Hasil"], SCHEDULE, [3.0, 2.0, 5.4, 3.2, 3.6])
    doc.add_heading("Sebelum bengkel", 1)
    for x in ["Uji akaun AI, Office/Google Workspace, Forms dan aliran automasi.", "Sediakan set data, fakta dokumen dan proses makmal yang sintetik atau diluluskan.", "Sediakan folder latihan, fail luar talian, tangkap layar dan output contoh.", "Pastikan peserta mengetahui larangan memasukkan data sensitif ke alat tidak diluluskan."]:
        bullet(doc, x)
    for idx, (name, duration, tools, activities, output) in enumerate(MODULES):
        doc.add_page_break()
        doc.add_heading(name, 1)
        p = doc.add_paragraph()
        r = p.add_run(f"Tempoh: {duration}  |  Alat: {tools}")
        set_font(r, size=10, bold=True, color=TEAL)
        doc.add_heading("Urutan penyampaian", 2)
        for activity in activities:
            bullet(doc, activity)
        doc.add_heading("Hasil peserta", 2)
        doc.add_paragraph(output)
        doc.add_heading("Cue fasilitator", 2)
        cues = ["Mulakan dengan masalah kerja dan bukti, bukan aplikasi.", "Tunjukkan satu contoh lengkap sebelum peserta mencuba.", "Minta peserta menyatakan sumber, batas dan semakan manusia.", "Semak hasil terhadap kriteria; jangan menilai berdasarkan rupa sahaja."]
        if idx == 2 or idx == 3:
            cues.append("Jangan benarkan AI mentafsir data sebelum jenis, unit dan kualitinya disahkan.")
        if idx == 4:
            cues.append("Automasi mesti mempunyai pemilik, laluan pengecualian dan syarat berhenti.")
        for cue in cues:
            bullet(doc, cue)
    doc.add_page_break()
    doc.add_heading("Pelan pemulihan", 1)
    table(doc, ["Gangguan", "Tindakan"], [
        ("AI tidak tersedia", "Gunakan contoh input/output bercetak dan lakukan audit manual."),
        ("Akaun Office/Google gagal", "Kerja berpasangan atau gunakan templat luar talian."),
        ("Automasi tidak berjalan", "Simulasikan borang, jadual status, kelulusan dan pemberitahuan."),
        ("Masa lewat", "Kekalkan hasil modul; kurangkan contoh dan perkongsian kumpulan."),
    ], [5.0, 12.2])
    doc.save(ROOT / "02_Facilitator-Guide" / "facilitator-guide.docx")


def response_box(doc, lines=3):
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    c = t.cell(0, 0)
    c.width = Cm(17.0)
    t.rows[0].height = Cm(max(2.0, lines * 0.9))
    t.rows[0].height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
    shade(c, "F8FAFA")
    margins(c, top=160, start=160, bottom=160, end=160)
    p = c.paragraphs[0]
    r = p.add_run("Catatan peserta")
    set_font(r, size=8.5, italic=True, color="7A8790")
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def build_workbook():
    doc = Document()
    configure(doc, "Buku Kerja Peserta — Bengkel Transformasi Digital MPE")
    title_block(doc, "Buku Kerja Peserta", "Transformasi Digital dan Pemerkasaan Operasi Pintar Makmal Penyelidikan Elektrik (MPE)\n28–29 Julai 2026")
    doc.add_heading("Hasil dua hari", 1)
    for x in ["Prompt produktiviti yang boleh diuji", "Draf dokumen rasmi dan semakan kelulusan", "Analisis data serta daftar rekod teknikal", "Peta proses dan prototaip aliran kerja", "Ringkasan projek, pembentangan dan pelan 30 hari"]:
        bullet(doc, x)
    doc.add_heading("Jadual kerja", 1)
    table(doc, ["Masa", "Modul", "Hasil"], [(a, b, e) for a, b, c, d, e in SCHEDULE], [4.0, 3.0, 10.2])
    prompts = [
        ("Modul 1 — Tugas dan prompt", ["Tugas berulang yang dipilih", "Pengguna, matlamat dan input", "Output, batas dan semakan manusia", "Bukti ujian: apa yang lulus atau gagal"]),
        ("Modul 2 — Dokumen rasmi", ["Tujuan, penerima dan tindakan yang diperlukan", "Fakta yang diluluskan", "Bahagian yang dijana atau dibantu AI", "Fakta, format dan kelulusan yang perlu disemak"]),
        ("Modul 3A — Data dan rekod", ["Soalan analisis", "Isu jenis data, unit, tarikh, nilai hilang dan pendua", "Formula/visual yang digunakan dan sebabnya", "Medan daftar rekod: ID, pemilik, status, versi, lokasi, tarikh"]),
        ("Modul 3B — Audit kualiti", ["Isu lengkap, unik, sah, konsisten dan terkini", "Pembetulan, pemilik dan jejak audit", "Dapatan yang disokong bukti", "Batas atau perkara yang belum diketahui"]),
        ("Modul 4 — Aliran kerja digital", ["Pencetus dan hasil proses", "Langkah, keputusan, menunggu dan pemilik semasa", "Validasi, kelulusan dan pengecualian masa hadapan", "Keputusan ujian laluan biasa, tidak lengkap dan eskalasi"]),
        ("Modul 5 — Projek dan pembentangan", ["Konteks dan masalah", "Bukti utama", "Cadangan dan implikasi", "Tindakan 30 hari, pemilik, ukuran dan syarat berhenti"]),
    ]
    for title, questions in prompts:
        doc.add_page_break()
        doc.add_heading(title, 1)
        for q in questions:
            doc.add_heading(q, 2)
            response_box(doc, 2)
    doc.add_page_break()
    doc.add_heading("Semakan penggunaan bertanggungjawab", 1)
    for x in ["Data sintetik, awam atau diluluskan sahaja", "Sumber dan bukti dapat dikenal pasti", "Maklumat hilang tidak direka", "Akses dan tindakan berada dalam skop", "Dokumen atau keputusan penting disemak manusia", "Terdapat laluan eskalasi dan syarat berhenti"]:
        bullet(doc, "☐ " + x)
    doc.add_heading("Komitmen 30 hari", 1)
    response_box(doc, 5)
    doc.save(ROOT / "03_Participant-Workbook" / "participant-workbook.docx")


if __name__ == "__main__":
    build_facilitator()
    build_workbook()
